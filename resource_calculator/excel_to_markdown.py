#!/usr/bin/env python3
"""
Excel to Markdown Converter with Formula Extraction
Converts each sheet in an Excel file to CSV format and combines them into a single Markdown file.
Also extracts and displays formulas used in each sheet.
"""

import pandas as pd
import openpyxl
from openpyxl import load_workbook
import csv
import io
import os
from datetime import datetime

def extract_formulas_from_sheet(workbook, sheet_name):
    """
    Extract formulas from a specific sheet in the workbook.
    Returns a dictionary with cell coordinates as keys and formulas as values.
    """
    formulas = {}
    sheet = workbook[sheet_name]
    
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == 'f':  # Formula cell
                try:
                    # Handle different types of formulas
                    if hasattr(cell, 'value') and cell.value is not None:
                        if isinstance(cell.value, str):
                            # Regular formula
                            formulas[f"{cell.coordinate}"] = cell.value
                        elif hasattr(cell.value, 'text'):
                            # Array formula or other complex formula
                            formulas[f"{cell.coordinate}"] = cell.value.text
                        else:
                            # Try to convert to string
                            formulas[f"{cell.coordinate}"] = str(cell.value)
                except Exception as e:
                    # If we can't extract the formula, note it
                    formulas[f"{cell.coordinate}"] = f"[Formula extraction error: {str(e)}]"
    
    return formulas

def sheet_to_csv_string(sheet_data):
    """
    Convert a pandas DataFrame to CSV string format.
    """
    csv_buffer = io.StringIO()
    sheet_data.to_csv(csv_buffer, index=False)
    return csv_buffer.getvalue()

def excel_to_markdown(excel_file_path, output_file_path=None):
    """
    Convert Excel file to Markdown format with CSV data and formulas.
    
    Args:
        excel_file_path (str): Path to the Excel file
        output_file_path (str): Path for the output Markdown file (optional)
    """
    
    if output_file_path is None:
        base_name = os.path.splitext(excel_file_path)[0]
        output_file_path = f"{base_name}_converted.md"
    
    # Load the Excel file with openpyxl for formula extraction
    workbook = load_workbook(excel_file_path, data_only=False)
    
    # Load with pandas for data extraction
    excel_file = pd.ExcelFile(excel_file_path)
    
    # Create the markdown content
    markdown_content = []
    
    # Add header
    markdown_content.append("# Excel to Markdown Conversion")
    markdown_content.append(f"**Source File:** {os.path.basename(excel_file_path)}")
    markdown_content.append(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    markdown_content.append(f"**Total Sheets:** {len(excel_file.sheet_names)}")
    markdown_content.append("")
    
    # Process each sheet
    for sheet_name in excel_file.sheet_names:
        print(f"Processing sheet: {sheet_name}")
        
        # Add sheet header
        markdown_content.append(f"## Sheet: {sheet_name}")
        markdown_content.append("")
        
        try:
            # Read sheet data
            sheet_data = pd.read_excel(excel_file_path, sheet_name=sheet_name)
            
            # Skip empty sheets
            if sheet_data.empty:
                markdown_content.append("*This sheet is empty.*")
                markdown_content.append("")
                continue
            
            # Convert to CSV and add to markdown
            csv_content = sheet_to_csv_string(sheet_data)
            markdown_content.append("### CSV Data")
            markdown_content.append("```csv")
            markdown_content.append(csv_content)
            markdown_content.append("```")
            markdown_content.append("")
            
            # Extract and add formulas
            formulas = extract_formulas_from_sheet(workbook, sheet_name)
            
            if formulas:
                markdown_content.append("### Formulas Used")
                markdown_content.append("")
                for cell_ref, formula in formulas.items():
                    markdown_content.append(f"- **{cell_ref}:** `{formula}`")
                markdown_content.append("")
            else:
                markdown_content.append("*No formulas found in this sheet.*")
                markdown_content.append("")
            
            # Add sheet separator
            markdown_content.append("---")
            markdown_content.append("")
            
        except Exception as e:
            markdown_content.append(f"*Error processing sheet: {str(e)}*")
            markdown_content.append("")
            markdown_content.append("---")
            markdown_content.append("")
    
    # Write to file
    with open(output_file_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(markdown_content))
    
    print(f"Conversion completed! Output saved to: {output_file_path}")
    return output_file_path

def main():
    """Main function to run the conversion."""
    excel_file = "Resource Calculator_V2_08SEP2025.xlsx"
    
    if not os.path.exists(excel_file):
        print(f"Error: Excel file '{excel_file}' not found!")
        return
    
    try:
        output_file = excel_to_markdown(excel_file)
        print(f"\nSuccessfully converted Excel file to Markdown!")
        print(f"Output file: {output_file}")
        
        # Display file size
        file_size = os.path.getsize(output_file)
        print(f"Output file size: {file_size:,} bytes")
        
    except Exception as e:
        print(f"Error during conversion: {str(e)}")

if __name__ == "__main__":
    main()
